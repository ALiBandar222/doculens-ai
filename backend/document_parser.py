
from io import BytesIO
from pathlib import Path
import re

import pymupdf
import pytesseract
from PIL import Image
from docx import Document
from pptx import Presentation


# Explicit path because Windows may not have Tesseract in PATH.
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


def _needs_arabic_ocr(text: str) -> bool:
    """Detect empty or suspiciously extracted PDF text."""

    if not text.strip():
        return True

    if "\ufffd" in text or "\u263b" in text:
        return True

    arabic_count = sum(
        "\u0600" <= char <= "\u06ff"
        for char in text
    )

    if arabic_count < 30:
        return False

    tokens = re.findall(
        r"[\u0600-\u06ff]+|[A-Za-z0-9]+",
        text,
    )

    if not tokens:
        return True

    single_character_ratio = (
        sum(len(token) == 1 for token in tokens) / len(tokens)
    )

    return single_character_ratio > 0.18



def _ocr_pdf_page(page) -> str:
    """Recognize Arabic PDF text using improved OCR settings."""

    pixmap = page.get_pixmap(dpi=400, alpha=False)
    image = Image.open(BytesIO(pixmap.tobytes("png"))).convert("L")

    from PIL import ImageOps
    image = ImageOps.autocontrast(image)

    return pytesseract.image_to_string(
        image,
        lang="ara",
        config="--psm 3",
    ).strip()


def extract_document_sections(
    filename: str,
    file_bytes: bytes,
) -> list[dict]:
    """
    Extract text from PDF, DOCX, PPTX, or TXT.

    PDFs retain page numbers. PowerPoint files retain slide numbers.
    Arabic OCR is used when PDF text appears unreliable.
    """

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    extension = Path(filename).suffix.lower()
    sections = []

    if extension == ".pdf":
        with pymupdf.open(
            stream=file_bytes,
            filetype="pdf",
        ) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()

                if _needs_arabic_ocr(text):
                    try:
                        ocr_text = _ocr_pdf_page(page)

                        if ocr_text:
                            text = ocr_text
                    except Exception as exc:
                        raise ValueError(
                            "PDF OCR failed. Check Tesseract "
                            "and its Arabic language data."
                        ) from exc

                if text:
                    sections.append({
                        "text": text,
                        "page_number": page_number,
                        "location_label": f"Page {page_number}",
                    })

    elif extension == ".docx":
        document = Document(BytesIO(file_bytes))
        content = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                content.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    content.append(" | ".join(cells))

        text = "\n".join(content).strip()

        if text:
            sections.append({
                "text": text,
                "page_number": None,
                "location_label": "Word document",
            })

    elif extension == ".pptx":
        presentation = Presentation(BytesIO(file_bytes))

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            content = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()

                    if text:
                        content.append(text)

                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]

                        if cells:
                            content.append(" | ".join(cells))

            text = "\n".join(content).strip()

            if text:
                sections.append({
                    "text": text,
                    "page_number": slide_number,
                    "location_label": f"Slide {slide_number}",
                })

    elif extension == ".txt":
        text = file_bytes.decode("utf-8-sig").strip()

        if text:
            sections.append({
                "text": text,
                "page_number": None,
                "location_label": "Text file",
            })

    else:
        raise ValueError(
            f"Unsupported file type: {extension or 'unknown'}"
        )

    if not sections:
        raise ValueError(
            "No readable text was found in this document."
        )

    return sections
